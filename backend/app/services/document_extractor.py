from __future__ import annotations
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import csv
from pypdf import PdfReader
from docx import Document

@dataclass
class ExtractedDocument:
    filename: str
    content_type: str
    text: str
    pages: int
    source_refs: list[str]

def _clean(text: str) -> str:
    return "\n".join(" ".join(x.split()) for x in text.replace("\x00", " ").splitlines() if x.strip())

def extract_document(filename: str, data: bytes) -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(BytesIO(data)); chunks=[]; refs=[]
        for n,page in enumerate(reader.pages,1):
            text=_clean(page.extract_text() or "")
            if text:
                chunks.append(f"[SOURCE: {filename} | PAGE {n}]\n{text}"); refs.append(f"{filename}#page={n}")
        if not chunks: raise ValueError("The PDF contains no extractable text. It may be scanned/image-only; OCR is required for this file.")
        return ExtractedDocument(filename,"application/pdf","\n\n".join(chunks),len(reader.pages),refs)
    if suffix == ".docx":
        doc=Document(BytesIO(data)); parts=[]; refs=[]
        for n,p in enumerate((p for p in doc.paragraphs if p.text.strip()),1):
            parts.append(f"[SOURCE: {filename} | PARAGRAPH {n}]\n{p.text.strip()}"); refs.append(f"{filename}#paragraph={n}")
        for ti,table in enumerate(doc.tables,1):
            rows=[" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
            if rows: parts.append(f"[SOURCE: {filename} | TABLE {ti}]\n"+"\n".join(rows)); refs.append(f"{filename}#table={ti}")
        if not parts: raise ValueError("The DOCX contains no readable text.")
        return ExtractedDocument(filename,"application/vnd.openxmlformats-officedocument.wordprocessingml.document","\n\n".join(parts),1,refs)
    if suffix in {".txt",".md",".markdown"}:
        text=_clean(data.decode("utf-8-sig",errors="replace"))
        if not text: raise ValueError("The text document is empty.")
        return ExtractedDocument(filename,"text/plain",f"[SOURCE: {filename} | TEXT]\n{text}",1,[f"{filename}#text"])
    if suffix == ".csv":
        rows=list(csv.reader(data.decode("utf-8-sig",errors="replace").splitlines()))
        if not rows: raise ValueError("The CSV is empty.")
        return ExtractedDocument(filename,"text/csv",f"[SOURCE: {filename} | CSV]\n"+"\n".join(" | ".join(r) for r in rows),1,[f"{filename}#csv"])
    raise ValueError("Unsupported file type. Upload PDF, DOCX, TXT, MD, or CSV.")
